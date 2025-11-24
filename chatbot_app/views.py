from io import BytesIO
import PyPDF2
import requests
import io
import json
import os
import traceback
import requests
from bs4 import BeautifulSoup
import csv
from django.http import JsonResponse, Http404
from django.conf import settings
from django.contrib.auth.models import User
from django.shortcuts import redirect, get_object_or_404
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from dotenv import load_dotenv
from openai import OpenAI
import openpyxl
from rest_framework.decorators import api_view, parser_classes
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from .models import AIModelConfig, PromptTemplate, JobDescriptionRevision, AIAuditLog, PuestoSAP, Administrador
from .models import JobDescription, JobDescriptionRevision
from .models import InformeSAP
from .serializers import InformeSAPSerializer
from rest_framework import status

from chatbot_app.models import ChatMessage, ChatSession, ChatErrorLog, TokenUsage
from .serializers import PuestoSAPSerializer
import pdfplumber
import pandas as pd
from django.http import JsonResponse
from chatbot_app.models import Candidato, ComparacionCVPuesto

# client = OpenAI()
load_dotenv()

client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))


def ejecutar_llamada_ia(funcion, variables: dict = {}):
    # Obtener modelo configurado
    config = AIModelConfig.objects.get(function=funcion)
    # Obtener template activo
    template = PromptTemplate.objects.filter(function=funcion, active=True).first()
    # Construir prompt dinámico
    prompt_texto = template.template.format(**variables)

    # Llamada al modelo
    response = client.chat.completions.create(
        model=config.model_name,
        messages=[
            {"role": "system", "content": config.description or "Actúa como asistente especializado."},
            {"role": "user", "content": prompt_texto}
        ]
    )

    return response, prompt_texto, config.model_name, template


# URLs autorizadas
AUTHORIZED_URLS = [
    "https://giintapeinnovahue.com/",
    "https://giintapeinnovahue.com/about.html"
    # "https://www.facebook.com/innsolcorporation",
]


# Función para extraer texto de URLs
def fetch_company_info():
    combined_text = []
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                      "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.5",
        "Accept-Encoding": "gzip, deflate",
        "Connection": "keep-alive",
    }
    for url in AUTHORIZED_URLS:
        try:
            resp = requests.get(url, headers=headers, timeout=5)
            resp.raise_for_status()
            soup = BeautifulSoup(resp.text, "html.parser")
            for script in soup(["script", "style"]):
                script.decompose()
            text = soup.get_text(separator=" ", strip=True)
            combined_text.append(text)
        except Exception as e:
            print(f"Error al extraer {url}: {e}")
    return " ".join(combined_text)


# Extraer y guardar contenido autorizado al cargar el backend
COMPANY_CONTEXT = fetch_company_info()[:1000]  # limitar tokens


# Función para llamar a OpenAI usando el contenido preprocesado
def call_openai_with_context(user_message):
    system_message = {
        "role": "system",
        "content": (
            "Eres un asistente virtual que responde únicamente sobre la empresa Giintape Innovahue. "
            f"Usa exclusivamente la información provista a continuación:\n{COMPANY_CONTEXT}\n\n"
            "Si la pregunta está fuera de este tema, responde exactamente: "
            "'Actualmente solo puedo responder dudas sobre Giintape Innovahue.'"
        )
    }

    try:
        # Llamada segura a OpenAI
        response = client.chat.completions.create(
            model="gpt-5-nano",
            messages=[
                system_message,
                {"role": "user", "content": user_message}
            ]
        )

        # Validación por si OpenAI no devuelve algo usable
        if not response or not response.choices:
            raise ValueError("La API no regresó una respuesta válida.")

        return response

    except Exception as e:
        # Guardar log del error
        ChatErrorLog.objects.create(
            user=None,  # O asigna request.user si lo usas dentro de un view
            session_id=None,  # O asigna la sesión actual si la tienes
            user_message=user_message,
            error_code=500,
            error_text=str(e)
        )

        # Re-lanzar para que el view maneje el error si es necesario
        raise e


# Decorador para requerir tokens
def require_tokens(func):
    def wrapper(request, *args, **kwargs):
        token_record = get_or_create_token_record(request)
        if not token_record.can_use_api():
            return JsonResponse({
                'error': 'Se alcanzó el límite mensual de tokens. Espera al siguiente ciclo.'
            }, status=403)
        return func(request, token_record=token_record, *args, **kwargs)

    return wrapper


@csrf_exempt
@require_tokens
def send_message(request, token_record):
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    try:
        data = json.loads(request.body)
        user_message = data.get('user_message', '').strip()
        if not user_message:
            return JsonResponse({'error': 'El mensaje está vacío'}, status=400)

        # Llamada a OpenAI
        '''
        response = client.chat.completions.create(
            model="gpt-5-nano",
            messages=[{"role": "user", "content": user_message}]
        )
        '''
        response = call_openai_with_context(user_message)

        bot_response = response.choices[0].message.content

        session = ChatSession.objects.filter(user=request.user).last()

        ChatMessage.objects.create(
            session=session,
            user_message=user_message,
            bot_response=bot_response
        )

        # Registrar tokens
        usage = getattr(response, 'usage', None)
        if usage:
            token_record.input_tokens += getattr(usage, 'prompt_tokens', 0)
            token_record.output_tokens += getattr(usage, 'completion_tokens', 0)
            token_record.save()

        return JsonResponse({
            'bot_response': bot_response,
            'tokens_restantes': {
                'entrada': token_record.remaining_input(),
                'salida': token_record.remaining_output(),
                'memoria': token_record.remaining_memory()
            }
        })

    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'error': str(e)}, status=500)


@csrf_exempt
def start_chat(request):
    # Crear una nueva sesión
    session = ChatSession.objects.create(user=request.user if request.user.is_authenticated else None)
    return redirect('chat', session_id=session.id)


@csrf_exempt
def chat(request, session_id):
    session = get_object_or_404(ChatSession, id=session_id)
    messages = list(session.messages.values(
        'id', 'user_message', 'bot_response', 'created_at'
    ))
    return JsonResponse({
        'session_id': session.id,
        'messages': messages
    })


'''
@csrf_exempt
def send_message_ajax(request, session_id):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            user_message = data.get('message')

            if not user_message:
                return JsonResponse({'error': 'No se recibió mensaje'}, status=400)

            response = client.chat.completions.create(
                model="gpt-5-nano",
                messages=[{"role": "user", "content": user_message}],
            )

            bot_response = response.choices[0].message.content

            session = get_object_or_404(ChatSession, id=session_id)

            chat_message = ChatMessage.objects.create(
                session=session,
                user_message=user_message,
                bot_response=bot_response,
                created_at=timezone.now()
            )

            return JsonResponse({
                'reply': bot_response,
                'user_message': user_message,
                'created_at': chat_message.created_at.strftime("%d %b %Y %H:%M")
            })

        except Exception as e:
            return JsonResponse({'error': str(e)}, status=500)

    return JsonResponse({'error': 'Método no permitido'}, status=405)
'''


@csrf_exempt
@require_tokens
def send_message_ajax(request, session_id, token_record):
    if request.method != "POST":
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    try:
        data = json.loads(request.body)
        user_message = data.get('message')
        if not user_message:
            return JsonResponse({'error': 'No se recibió mensaje'}, status=400)

        session = get_object_or_404(ChatSession, id=session_id)
        token_record.session = session
        token_record.save(update_fields=['session'])

        # Llamada a OpenAI
        '''
                response = client.chat.completions.create(
                    model="gpt-5-nano",
                    messages=[{"role": "user", "content": user_message}]
                )
                '''
        response = call_openai_with_context(user_message)

        bot_response = response.choices[0].message.content

        session = ChatSession.objects.filter(user=request.user).last()

        chat_message = ChatMessage.objects.create(
            session=session,
            user_message=user_message,
            bot_response=bot_response
        )

        # Registrar tokens
        usage = getattr(response, 'usage', None)
        if usage:
            token_record.input_tokens += getattr(usage, 'prompt_tokens', 0)
            token_record.output_tokens += getattr(usage, 'completion_tokens', 0)
            token_record.save()

        return JsonResponse({
            'reply': bot_response,
            'user_message': user_message,
            'created_at': chat_message.created_at.strftime("%d %b %Y %H:%M"),
            'tokens_restantes': {
                'entrada': token_record.remaining_input(),
                'salida': token_record.remaining_output(),
                'memoria': token_record.remaining_memory()
            }
        })

    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'error': str(e)}, status=500)


@csrf_exempt
def log_error(request):
    if request.method == "POST":
        data = json.loads(request.body)
        ChatErrorLog.objects.create(
            user=request.user if request.user.is_authenticated else None,
            session_id=data.get("session_id"),
            user_message=data.get("user_message"),
            error_code=data.get("error_code"),
            error_text=data.get("error_text")
        )
        return JsonResponse({"status": "ok"})
    return JsonResponse({"error": "Método no permitido"}, status=405)


def ver_errores(request):
    errores = ChatErrorLog.objects.all().order_by('-created_at').values(
        'id', 'user__username', 'session_id', 'user_message', 'error_text', 'error_code', 'created_at'
    )
    return JsonResponse(list(errores), safe=False)


def exportar_errores_excel(request):
    # Obtener los errores
    errores = ChatErrorLog.objects.all().order_by('-created_at').values(
        'id', 'user__username', 'session_id', 'user_message', 'error_text', 'error_code', 'created_at'
    )

    # Crear un libro y una hoja
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Errores API"

    # Encabezados
    headers = ['ID', 'Usuario', 'Session ID', 'Mensaje', 'Error', 'Código', 'Fecha']
    ws.append(headers)

    # Agregar los datos
    for err in errores:
        ws.append([
            err['id'],
            err['user__username'] or 'Anon',
            err['session_id'] or '-',
            err['user_message'],
            err['error_text'],
            err['error_code'] or '-',
            err['created_at'].strftime("%Y-%m-%d %H:%M:%S")
        ])

    # Ajustar ancho de columnas automáticamente
    for col in ws.columns:
        max_length = max(len(str(cell.value)) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_length + 2

    # Preparar respuesta HTTP
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=errores_api.xlsx'
    wb.save(response)
    return response


@csrf_exempt
@require_tokens
def mejorar_descripcion_puesto(request, token_record):
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    try:
        data = json.loads(request.body)

        candidato_info = data.get('candidato', {})
        puesto_actual = data.get('puesto_actual', '')

        if not puesto_actual or not candidato_info:
            return JsonResponse({'error': 'Datos incompletos'}, status=400)

        # Llamar a IA con plantilla desde BD
        response, prompt_usado, modelo, template = ejecutar_llamada_ia(
            funcion="mejorar_descripcion",
            variables={
                "puesto": puesto_actual,
                "candidato": candidato_info.get("nombre", "")
            }
        )

        descripcion_mejorada = response.choices[0].message.content

        return JsonResponse({
            "descripcion_mejorada": descripcion_mejorada,
            "puesto_original": puesto_actual,
            "candidato": candidato_info.get("nombre", "")
        })

    except Exception as e:
        traceback.print_exc()
        return JsonResponse({"error": str(e)}, status=500)


def token_usage_list(request):
    """
    Lista todos los registros de tokens del mes actual.
    Para administradores o panel general.
    """
    now = timezone.now()
    month_start = now.replace(day=1)
    tokens = TokenUsage.objects.filter(month=month_start).select_related('user')

    # Devolver JSON simple
    data = [
        {
            'user': t.user.username if t.user else 'Anon',
            'input_tokens': t.input_tokens,
            'output_tokens': t.output_tokens,
            'memory_tokens': t.memory_tokens,
            'remaining_input': t.remaining_input(),
            'remaining_output': t.remaining_output(),
            'remaining_memory': t.remaining_memory()
        }
        for t in tokens
    ]

    # 🔍 Mostrar en la consola del servidor Django
    print("[API] Token usage data:")
    for d in data:
        print(d)

    return JsonResponse({'data': data})


def token_usage_detail(request, user_id):
    """
    Ver los tokens de un usuario específico del mes actual.
    """
    now = timezone.now()
    month_start = now.replace(day=1)
    user = get_object_or_404(User, id=user_id)
    token_record = TokenUsage.objects.filter(user=user, month=month_start).first()

    if not token_record:
        return JsonResponse({'error': 'No se encontraron registros de tokens para este usuario'}, status=404)

    data = {
        'user': user.username,
        'input_tokens': token_record.input_tokens,
        'output_tokens': token_record.output_tokens,
        'memory_tokens': token_record.memory_tokens,
        'remaining_input': token_record.remaining_input(),
        'remaining_output': token_record.remaining_output(),
        'remaining_memory': token_record.remaining_memory()
    }
    return JsonResponse(data)


def get_or_create_token_record(request):
    now = timezone.now().date()
    month_start = now.replace(day=1)

    user = request.user if request.user.is_authenticated else None
    session_key = request.session.session_key

    if user:
        token_record, created = TokenUsage.objects.get_or_create(
            user=user,
            month=month_start
        )

    else:
        # Nuevo: usar anon_id para anónimos
        anon_id = session_key or f"anon_{request.META.get('REMOTE_ADDR', 'unknown')}"

        token_record, created = TokenUsage.objects.get_or_create(
            anon_id=anon_id,
            month=month_start
        )

    return token_record


@csrf_exempt
def check_token_status(request):
    """
    Vista para que el frontend verifique el estado de los tokens
    """
    token_record = get_or_create_token_record(request)

    # Calcular próximo ciclo
    next_month = token_record.month.month % 12 + 1
    next_year = token_record.month.year + (1 if next_month == 1 else 0)
    next_cycle = timezone.datetime(next_year, next_month, 1).date()

    return JsonResponse({
        'can_use_api': token_record.can_use_api(),
        'is_blocked': token_record.is_blocked(),
        'input_tokens': token_record.input_tokens,
        'output_tokens': token_record.output_tokens,
        'remaining_input': token_record.remaining_input(),
        'remaining_output': token_record.remaining_output(),
        'max_tokens': token_record.MAX_TOKENS,
        'notified_limit': token_record.notified_limit,
        'current_cycle': token_record.month.strftime("%Y-%m"),
        'next_reset_date': next_cycle.strftime("%Y-%m-%d")
    })


# FUNCIÓN: Cambiar modelo IA desde panel
@csrf_exempt
def admin_set_model(request):
    """
    Cambia el modelo asignado a una función IA.
    Ejemplo POST:
    {
      "function": "mejorar_descripcion",
      "model_name": "gpt-5-nano"
    }
    """
    if request.method != "POST":
        return JsonResponse({"error": "Método no permitido"}, status=405)

    data = json.loads(request.body)
    function = data.get("function")
    model_name = data.get("model_name")

    config = AIModelConfig.objects.filter(function=function).first()
    if not config:
        return JsonResponse({"error": "Función no registrada"}, status=404)

    config.model_name = model_name
    config.save()

    return JsonResponse({
        "status": "ok",
        "message": f"Modelo actualizado para {function}",
        "model": model_name
    })


# FUNCIÓN: Crear nuevo template
@csrf_exempt
def admin_create_prompt(request):
    """
    Crear un template nuevo.
    POST:
    {
        "name": "template_reclutamiento_v2",
        "function": "mejorar_descripcion",
        "template": "Mejora el puesto {puesto}..."
    }
    """
    if request.method != "POST":
        return JsonResponse({"error": "Método no permitido"}, status=405)

    data = json.loads(request.body)

    template = PromptTemplate.objects.create(
        name=data["name"],
        function=data.get("function", ""),
        template=data["template"],
        active=True
    )

    return JsonResponse({
        "status": "ok",
        "id": template.id,
        "message": "Template creado"
    })


# FUNCIÓN: Activar un template (desactivar los demás)
@csrf_exempt
def admin_activate_prompt(request):
    """
    Activa un template y desactiva el resto.
    POST: { "template_id": 5 }
    """
    data = json.loads(request.body)
    template_id = data.get("template_id")

    template = PromptTemplate.objects.filter(id=template_id).first()
    if not template:
        return JsonResponse({"error": "Template no encontrado"}, status=404)

    PromptTemplate.objects.filter(function=template.function).update(active=False)

    template.active = True
    template.save()

    return JsonResponse({
        "status": "ok",
        "activated_template": template.name
    })


# FUNCIÓN: Registrar revisión generada (cuando la IA devuelve respuesta)
def registrar_revision_puesto(job_description, response, prompt, modelo, admin, template, usage=None):
    """
    Guarda la versión generada en JobDescriptionRevision.
    """
    revision = JobDescriptionRevision.objects.create(
        job_description=job_description,
        texto_mejorado=response,
        prompt_usado=prompt,
        modelo_usado=modelo,
        prompt_template=template,
        creado_por=admin,
        version=(job_description.revisions.count() + 1),
        input_tokens=getattr(usage, "prompt_tokens", 0) if usage else 0,
        output_tokens=getattr(usage, "completion_tokens", 0) if usage else 0
    )

    job_description.latest_revision = revision
    job_description.save()

    return revision


# FUNCIÓN: Publicar una revisión final
@csrf_exempt
def admin_publicar_revision(request):
    """
    POST: { "revision_id": 10 }
    """
    data = json.loads(request.body)
    rev_id = data.get("revision_id")

    revision = JobDescriptionRevision.objects.filter(id=rev_id).first()
    if not revision:
        return JsonResponse({"error": "Revisión no encontrada"}, status=404)

    revision.is_published = True
    revision.job_description.publicado = True

    revision.save()
    revision.job_description.save()

    return JsonResponse({"status": "ok", "message": "Revisión publicada"})


# FUNCIÓN: Guardar auditoría completa
def registrar_auditoria(revision, request_payload, response_payload, error=None):
    return AIAuditLog.objects.create(
        revision=revision,
        request_payload=request_payload,
        response_payload=response_payload,
        error_text=error
    )


@csrf_exempt
def prompt_list_view(request):
    """
    GET: /prompt/list/
    Devuelve { data: [ {id,name,function,template,active,created_at} ] }
    """
    if request.method != 'GET':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    prompts = PromptTemplate.objects.all().order_by('-created_at').values(
        'id', 'name', 'function', 'template', 'active', 'created_at'
    )
    data = []
    for p in prompts:
        data.append({
            'id': p['id'],
            'name': p['name'],
            'function': p['function'],
            'template': p['template'],
            'active': bool(p['active']),
            'created_at': p['created_at'].strftime("%Y-%m-%d %H:%M:%S") if p['created_at'] else None
        })
    return JsonResponse({'data': data})


@csrf_exempt
def revisions_list_view(request):
    """
    GET: /revisions/list/
    Devuelve { data: [ {id, job_description_id, version, modelo_usado, creado_en, texto_mejorado} ] }
    """
    if request.method != 'GET':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    revs = JobDescriptionRevision.objects.select_related('job_description').all().order_by('-creado_en')[:200]
    data = []
    for r in revs:
        data.append({
            'id': r.id,
            'job_description_id': r.job_description.id if r.job_description else None,
            'version': r.version,
            'modelo_usado': r.modelo_usado,
            'creado_en': r.creado_en.strftime("%Y-%m-%d %H:%M:%S") if r.creado_en else None,
            'texto_mejorado': (r.texto_mejorado[:400] + '...') if r.texto_mejorado and len(
                r.texto_mejorado) > 400 else (r.texto_mejorado or '')
        })
    return JsonResponse({'data': data})


# informes generados de SAP SuccessFactors (SSFF), ingresados manualmente(en el cuadro de diálogo) o desde “Gestión de posiciones”

@api_view(['POST'])
def import_puesto_sap(request):
    serializer = PuestoSAPSerializer(data=request.data)

    if serializer.is_valid():
        serializer.save()
        return Response({"message": "Puesto importado correctamente."}, status=status.HTTP_201_CREATED)

    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(['GET'])
def listar_puestos(request):
    puestos = PuestoSAP.objects.all()
    serializer = PuestoSAPSerializer(puestos, many=True)
    return Response(serializer.data)


@api_view(['POST'])
@parser_classes([MultiPartParser])
def importar_puestos_csv(request):
    file = request.FILES.get("file")

    if file is None:
        return Response({"error": "No se recibió archivo"}, status=status.HTTP_400_BAD_REQUEST)

    decoded_file = file.read().decode("utf-8").splitlines()
    reader = csv.DictReader(decoded_file)

    registros = 0
    for row in reader:  # type: dict
        PuestoSAP.objects.update_or_create(
            id_requisicion=int(row["ID de requisición de personal"]),
            defaults={
                "categoria": row["Categoría"],
                "titulo": row["Titulo"],
                "link": row["link"],
                "ubicacion": row["Ubicación"],
            }
        )
        registros += 1

    return Response({"message": "Importación completada", "total": registros})


@api_view(['POST'])
def agregar_informe_sap(request):
    serializer = InformeSAPSerializer(data=request.data)
    if serializer.is_valid():
        serializer.save()
        return Response({"message": "Informe SAP agregado correctamente"}, status=status.HTTP_201_CREATED)
    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(['GET'])
def listar_informes_sap(request):
    informes = InformeSAP.objects.all()
    serializer = InformeSAPSerializer(informes, many=True)
    return Response(serializer.data)


@api_view(['POST'])
def importar_puestos_desde_urls(request):
    informes = InformeSAP.objects.all()
    total_registros = 0
    errores = []

    for informe in informes:
        try:
            r = requests.get(informe.url, timeout=10)
            r.raise_for_status()  # Lanza excepción si status != 200

            decoded_file = r.content.decode('utf-8').splitlines()
            reader = csv.DictReader(decoded_file)

            registros = 0
            for row in reader:  # type: dict
                # Convertir ID a entero
                try:
                    id_req = int(row.get("ID de requisición de personal", 0))
                except ValueError:
                    errores.append(f"Fila con ID inválido en informe {informe.nombre}: {row}")
                    continue

                if id_req == 0:
                    continue  # Saltar filas vacías o mal formateadas

                PuestoSAP.objects.update_or_create(
                    id_requisicion=id_req,
                    defaults={
                        "categoria": row.get("Categoría") or row.get("Categoria", ""),
                        "titulo": row.get("Titulo", ""),
                        "link": row.get("link", ""),
                        "ubicacion": row.get("Ubicación") or row.get("Ubicacion", ""),
                    }
                )
                registros += 1

            informe.ultima_descarga = timezone.now()
            informe.save()
            total_registros += registros

        except Exception as e:
            errores.append(f"Error al procesar {informe.url}: {str(e)}")

    return Response({
        "message": "Proceso completado",
        "total_registros_importados": total_registros,
        "errores": errores
    }, status=status.HTTP_200_OK)


ARCHIVOS_DIR = os.path.join(settings.BASE_DIR, 'archivos_sap')
os.makedirs(ARCHIVOS_DIR, exist_ok=True)

'''
@api_view(['POST'])
def importar_puestos_desde_urls(request):
    informes = InformeSAP.objects.all()
    total_registros = 0
    errores = []

    for informe in informes:
        try:
            r = requests.get(informe.url, timeout=10)
            r.raise_for_status()  # Lanza excepción si status != 200

            # Guardar archivo CSV
            fecha_str = datetime.now().strftime("%Y%m%d_%H%M%S")
            nombre_archivo = f"puestos_{fecha_str}.csv"
            ruta_archivo = os.path.join(ARCHIVOS_DIR, nombre_archivo)
            with open(ruta_archivo, "wb") as f:
                f.write(r.content)

            # Leer CSV en memoria
            decoded_file = r.content.decode('utf-8').splitlines()
            reader = csv.DictReader(decoded_file)

            registros = 0
            for row in reader:
                try:
                    id_req = int(row.get("ID de requisición de personal", 0))
                except ValueError:
                    errores.append(f"Fila con ID inválido en informe {informe.nombre}: {row}")
                    continue
                if id_req == 0:
                    continue

                PuestoSAP.objects.update_or_create(
                    id_requisicion=id_req,
                    defaults={
                        "categoria": row.get("Categoría") or row.get("Categoria", ""),
                        "titulo": row.get("Titulo", ""),
                        "link": row.get("link", ""),
                        "ubicacion": row.get("Ubicación") or row.get("Ubicacion", ""),
                    }
                )
                registros += 1

            informe.ultima_descarga = timezone.now()
            informe.save()
            total_registros += registros

        except Exception as e:
            errores.append(f"Error al procesar {informe.url}: {str(e)}")

    return Response({
        "message": "Proceso completado",
        "total_registros_importados": total_registros,
        "errores": errores,
        "archivos_guardados": [nombre_archivo]  # Lista de archivos guardados
    }, status=status.HTTP_200_OK)
'''


@csrf_exempt
def crear_perfil_manual(request):
    """
    Crea un perfil de puesto manual y su primera revisión.
    No depende de usuarios ni autenticación.
    """
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido'}, status=405)

    texto = request.POST.get('descripcion', '').strip()
    if not texto:
        return JsonResponse({'error': 'La descripción no puede estar vacía'}, status=400)

    # Crear perfil de puesto manual
    job = JobDescription.objects.create(
        texto_original=texto,
        fuente=JobDescription.SOURCE_MANUAL
    )

    # Crear primera revisión inicial
    JobDescriptionRevision.objects.create(
        job_description=job,
        texto_mejorado=texto,
        version=1
    )

    return JsonResponse({'success': True, 'job_id': job.id})


@csrf_exempt
def lista_perfiles(request):
    """
    Devuelve todos los perfiles de puesto como JSON.
    """
    perfiles = JobDescription.objects.all().order_by('-creado_en')
    data = []
    for perfil in perfiles:
        data.append({
            'id': perfil.id,
            'texto_original': perfil.texto_original,
            'fuente': perfil.fuente,
            'creado_en': perfil.creado_en.isoformat(),
            'publicado': perfil.publicado
        })

    return JsonResponse({'perfiles': data})


def ver_perfil(request, id):
    try:
        perfil = JobDescription.objects.select_related("latest_revision").get(id=id)
    except JobDescription.DoesNotExist:
        raise Http404("Perfil no encontrado")

    data = {
        "id": perfil.id,
        "texto_original": perfil.texto_original,
        "fuente": perfil.fuente,
        "creado_en": perfil.creado_en,
        "actualizado_en": perfil.actualizado_en,
        "publicado": perfil.publicado,
        "latest_revision": None
    }

    # Si tiene revisión
    if perfil.latest_revision:
        rev = perfil.latest_revision
        data["latest_revision"] = {
            "version": rev.version,
            "texto_mejorado": rev.texto_mejorado,
            "creado_en": rev.creado_en,
            "is_published": rev.is_published,
            "input_tokens": rev.input_tokens,
            "output_tokens": rev.output_tokens,
        }

    return JsonResponse(data, safe=False)


# COMPARACION DE CV´S

'''
def descargar_texto_cv(url_pdf: str) -> str:
    r = requests.get(url_pdf)
    path = "/tmp/cv.pdf"
    with open(path, "wb") as f:
        f.write(r.content)

    with pdfplumber.open(path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)

'''

import requests
import pdfplumber
from io import BytesIO


def descargar_texto_cv(url_pdf: str) -> str:
    print("=== [LOG] Intentando descargar PDF ===")
    print("[LOG] URL ORIGINAL:", url_pdf)

    # 1) Primer intento (quizá funciona)
    r = requests.get(url_pdf, timeout=15)

    # Si SharePoint devuelve HTML en vez de PDF, forzamos el download
    if "pdf" not in r.headers.get("Content-Type", "").lower():
        print("[WARN] SharePoint devolvió HTML, forzando descarga real del PDF...")

        # Forzar link directo
        if "?download=1" not in url_pdf:
            if "?" in url_pdf:
                url_pdf = url_pdf + "&download=1"
            else:
                url_pdf = url_pdf + "?download=1"

        print("[LOG] URL FORZADA:", url_pdf)

        r = requests.get(url_pdf, timeout=15)

    # Si aún no es PDF → NO sirve
    if "pdf" not in r.headers.get("Content-Type", "").lower():
        print("[ERROR] SharePoint sigue sin devolver PDF. Content-Type:", r.headers.get("Content-Type"))
        return "ERROR: NO_ES_PDF"

    print("[LOG] PDF DESCARGADO CORRECTAMENTE")

    # 2) Extraer texto del PDF
    try:
        pdf_bytes = BytesIO(r.content)
        with pdfplumber.open(pdf_bytes) as pdf:
            texto = "\n".join((p.extract_text() or "") for p in pdf.pages)
        print("[LOG] Texto extraído correctamente")
        return texto.strip()
    except Exception as e:
        print("[ERROR] NO_SE_PUDO_LEER_PDF:", e)
        return "ERROR: NO_SE_PUDO_LEER_PDF"


def cargar_puesto(req_id: str) -> str:
    df = pd.read_csv("puestos.csv", encoding="utf-8")

    fila = df[df["reqId_ix"].astype(str) == str(req_id)]

    if fila.empty:
        raise ValueError("Puesto no encontrado en CSV")

    texto = fila["jobDesc_ix"].iloc[0]

    if pd.isna(texto):
        raise ValueError("La descripción del puesto está vacía (NaN en CSV)")

    return str(texto)


def comparar_cv_con_puesto(candidato, req_id: str):
    # 1. Obtener URL del CV
    cv_url = candidato.CV_candidate
    if not cv_url:
        raise ValueError("El candidato no tiene CV en campo CV_candidate")

    # 2. Extraer texto del CV
    cv_texto = descargar_texto_cv(cv_url)

    # 3. Cargar descripción del puesto
    puesto_texto = cargar_puesto(req_id)

    # 4. Llamar a OpenAI
    prompt = f"""
Analiza el CV y compáralo con el puesto.

### CV
{cv_texto}

### PUESTO
{puesto_texto}

### RESPUESTA
Incluye JSON con:
- compatibilidad (0-100)
- fortalezas
- debilidades
- resumen
"""

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )

    resultado_texto = response.choices[0].message.content

    # Podríamos intentar extraer el número automáticamente, por ahora manual:
    score = 0
    for s in range(100, -1, -1):
        if str(s) in resultado_texto:
            score = s
            break

    # 5. Guardar en BD
    registro = ComparacionCVPuesto.objects.create(
        candidato=candidato,
        id_puesto=req_id,
        resultado=resultado_texto,
        score=score
    )

    return registro


def comparar_candidato_view(request, candidato_id, req_id):
    print("=== [LOG] Iniciando comparar_candidato_view ===")
    print(f"[LOG] candidato_id recibido: {candidato_id}")
    print(f"[LOG] req_id recibido: {req_id}")

    try:
        # 1. Obtener candidato
        print("[LOG] Buscando candidato...")
        candidato = Candidato.objects.get(id_candidate=candidato_id)
        print("[LOG] Candidato encontrado correctamente")
        print(f"[LOG] Datos del candidato: {candidato.nombre_candidate} {candidato.apellidop_candidate}")

        # 2. Cargar la descripción del puesto DESDE CSV
        print("[LOG] Cargando puesto desde CSV...")
        try:
            texto_req = cargar_puesto(req_id)
            print("[LOG] Puesto cargado correctamente")
            print("[LOG] Texto del puesto:\n", texto_req[:300], "...")
        except Exception as e:
            print("[ERROR] No se pudo cargar el puesto desde CSV:", e)
            return JsonResponse({"error": "Puesto no encontrado en CSV"}, status=404)

        # 3. Armar texto del candidato (solo para logs, no para IA)
        texto_candidato = f"""
        Nombre: {candidato.nombre_candidate} {candidato.apellidop_candidate} {candidato.apellidom_candidate}
        Correo: {candidato.correo_candidate}
        Tel: {candidato.tel_candidate}
        CV: {candidato.CV_candidate}
        """
        print("[LOG] Texto armado del candidato:")
        print(texto_candidato)

        # 4. Enviar a IA usando tu función REAL
        print("[LOG] Enviando a IA para comparación...")
        resultado = comparar_cv_con_puesto(candidato, req_id)
        print("[LOG] Respuesta de la IA recibida:")
        print(resultado)

        print("[LOG] Todo OK, regresando JSON")

        return JsonResponse({
            "ok": True,
            "req_id": req_id,
            "candidato_id": candidato_id,
            "resultado_id": resultado.id,
            "score": resultado.score,
        })

    except Candidato.DoesNotExist:
        print("[ERROR] Candidato no existe")
        return JsonResponse({"error": "Candidato no existe"}, status=404)

    except Exception as e:
        print("[ERROR] Excepción inesperada:")
        import traceback
        traceback.print_exc()
        print(f"[ERROR] {str(e)}")
        return JsonResponse({"error": str(e)}, status=500)


# Mejorar puesto de archivo PUESTO.CSV


def mejorar_puesto_y_guardar(req_id: str) -> str:
    import pandas as pd
    from django.utils import timezone
    from .models import Puesto

    # Convertir a entero para que coincida con el CSV
    try:
        req_int = int(req_id)
    except:
        raise ValueError("req_id inválido")

    # 1. Leer CSV
    df = pd.read_csv("puestos.csv", encoding="utf-8")

    # Asegurar que la columna es numérica
    df["reqId_ix"] = pd.to_numeric(df["reqId_ix"], errors="coerce")

    fila = df[df["reqId_ix"].astype(str) == str(req_id)]

    if fila.empty:
        raise ValueError(f"Puesto {req_id} NO está en el CSV")

    descripcion_original = fila["jobDesc_ix"].iloc[0]

    # 2. Llamar a la IA
    prompt = (
        "Mejora profesionalmente la siguiente descripción sin inventar funciones nuevas:\n\n"
        f"{descripcion_original}"
    )

    response = client.chat.completions.create(
        model="gpt-5-nano",
        messages=[
            {"role": "system", "content": "Eres experto en redacción corporativa."},
            {"role": "user", "content": prompt}
        ]
    )

    # descripcion_mejorada = response.choices[0].message["content"]
    descripcion_mejorada = response.choices[0].message.content

    # 3. Guardar en BD
    puesto, _ = Puesto.objects.update_or_create(
        req_id=req_int,
        defaults={
            "descripcion_original": descripcion_original,
            "descripcion_mejorada": descripcion_mejorada,
            "fecha_mejora": timezone.now()
        }
    )

    return descripcion_mejorada


# SELECCION WEB

'''
def mejorar_puesto_view(request, req_id):
    aplicar = request.GET.get("mejorar", "1") == "1"
    resultado = mejorar_descripcion_puesto_2(req_id, aplicar_mejora=aplicar)
    return JsonResponse(resultado, safe=False)
'''


def mejorar_puesto_view(request, req_id):
    try:
        nueva_desc = mejorar_puesto_y_guardar(req_id)
        return JsonResponse({"descripcion_mejorada": nueva_desc})
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


##  COMPARAR CV CON DESCRIPCIÓN MANUAL


def comparar_cv_con_texto_manual(candidato, texto_manual: str):
    """
    Compara el CV del candidato contra una descripción manual escrita en un textarea.
    """

    # 1. Obtener URL del CV
    cv_url = candidato.CV_candidate
    if not cv_url:
        raise ValueError("El candidato no tiene URL de CV")

    # 2. Extraer texto real del PDF del CV
    cv_texto = descargar_texto_cv(cv_url)

    # 3. Preparar prompt para IA
    prompt = f"""
Analiza el CV y compáralo con la descripción manual.

### CV
{cv_texto}

### DESCRIPCIÓN MANUAL
{texto_manual}

### RESPUESTA
Devuelve:
- compatibilidad (0-100)
- fortalezas
- debilidades
- resumen
En formato JSON.
"""

    # 4. Llamar a OpenAI
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )

    resultado_texto = response.choices[0].message.content

    # 5. Extraer score mínimo
    score = 0
    for s in range(100, -1, -1):
        if str(s) in resultado_texto:
            score = s
            break

    # 6. Guardar en DB
    registro = ComparacionCVPuesto.objects.create(
        candidato=candidato,
        id_puesto="manual",  # o NULL si tu modelo lo permite
        resultado=resultado_texto,
        score=score
    )

    return registro


@csrf_exempt
def comparar_candidato_manual_view(request, candidato_id):
    """
    Compara el CV de un candidato contra un texto manual escrito en un textarea.
    """
    print("=== [LOG] Iniciando comparar_candidato_manual_view ===")
    print(f"[LOG] candidato_id recibido: {candidato_id}")

    if request.method != "POST":
        return JsonResponse({"error": "POST requerido"}, status=400)

    # 1. Leer texto manual desde request
    data = json.loads(request.body.decode("utf-8"))
    texto_manual = data.get("texto_manual", "").strip()

    if not texto_manual:
        return JsonResponse({"error": "Se requiere texto_manual"}, status=400)

    print("[LOG] Texto manual recibido:")
    print(texto_manual[:200], "...")

    try:
        # 2. Obtener candidato de BD
        print("[LOG] Buscando candidato...")
        candidato = Candidato.objects.get(id_candidate=candidato_id)
        print("[LOG] Candidato encontrado correctamente")

        # 3. Realizar comparación
        print("[LOG] Comparando CV contra descripción manual...")
        resultado = comparar_cv_con_texto_manual(candidato, texto_manual)

        print("[LOG] Comparación completada correctamente")

        return JsonResponse({
            "ok": True,
            "candidato_id": candidato_id,
            "resultado_id": resultado.id,
            "score": resultado.score,
        })

    except Candidato.DoesNotExist:
        print("[ERROR] Candidato no existe")
        return JsonResponse({"error": "Candidato no existe"}, status=404)

    except Exception as e:
        print("[ERROR] Excepción inesperada:")
        import traceback
        traceback.print_exc()
        return JsonResponse({"error": str(e)}, status=500)


''' SOLO PARA DETALLES '''


@csrf_exempt
def extraer_puesto_candidato(request, candidato_id):
    """
    Extrae el puesto del CV del candidato para mostrar en el modal
    """
    try:
        candidato = Candidato.objects.get(id_candidate=candidato_id)

        puesto_extraido = "No especificado"

        if candidato.CV_candidate:
            puesto_extraido = extraer_puesto_desde_cv(candidato.CV_candidate)
            # Actualizar en la base de datos para futuras consultas
            candidato.puesto = puesto_extraido
            candidato.save()

        return JsonResponse({
            "puesto_extraido": puesto_extraido,
            "candidato_id": candidato_id
        })

    except Candidato.DoesNotExist:
        return JsonResponse({"error": "Candidato no existe"}, status=404)
    except Exception as e:
        return JsonResponse({"puesto_extraido": "No especificado"})  # Fallback


def extraer_puesto_desde_cv(cv_url):
    """
    Intenta extraer el puesto/posición del texto del CV PDF
    """
    try:
        # Descargar el PDF
        response = requests.get(cv_url)
        pdf_file = io.BytesIO(response.content)

        # Leer PDF
        pdf_reader = PyPDF2.PdfReader(pdf_file)
        texto_cv = ""

        for page in pdf_reader.pages:
            texto_cv += page.extract_text()

        # Buscar patrones comunes de puesto
        patrones_puesto = [
            "Objective:",
            "Summary:",
            "Position:",
            "Puesto:",
            "Cargo:",
            "Título:",
            "Applied for:",
            "Seeking:",
            "Aspirante a:"
        ]

        # Buscar en las primeras líneas (donde suele estar el puesto)
        lineas = texto_cv.split('\n')
        for i, linea in enumerate(lineas[:20]):  # Primeras 20 líneas
            linea_limpia = linea.strip()

            # Si la línea parece un título/puesto (no vacía, no demasiado larga)
            if (linea_limpia and
                    len(linea_limpia) < 100 and
                    any(palabra in linea_limpia for palabra in
                        ['Engineer', 'Developer', 'Manager', 'Analyst', 'Coordinator', 'Specialist'])):
                return linea_limpia

            # Buscar después de patrones clave
            for patron in patrones_puesto:
                if patron.lower() in linea.lower():
                    if i + 1 < len(lineas):
                        siguiente_linea = lineas[i + 1].strip()
                        if siguiente_linea:
                            return siguiente_linea

        return "Extraído del CV"  # Fallback

    except Exception as e:
        print(f"Error extrayendo puesto del CV: {e}")
        return "No especificado"

